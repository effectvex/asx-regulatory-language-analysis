import csv
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

root = Path(__file__).resolve().parents[1]
source = root / "output" / "regulatory-terms-2015-2025.xlsx"
annual = root / "data" / "graph-data" / "annual-incidence.csv"
output = root / "output" / "regulatory-terms-2015-2025-with-incidence-chart.xlsx"

wb = load_workbook(source)
if "Incidence per 100k" in wb.sheetnames:
    del wb["Incidence per 100k"]
ws = wb.create_sheet("Incidence per 100k")
headers = ["Year", "Companies", "Reports", "Extracted words", "Coded passages", "Six-term incidence per 100,000 words"]
ws.append(headers)
with annual.open(newline="", encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        ws.append([
            int(row["Year"]), int(row["Companies"]), int(row["Reports"]),
            int(float(row["Words"])), int(float(row["CodedPassages"])),
            float(row["SixTermIncidencePer100k"]),
        ])

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1769AA")
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
ws.freeze_panes = "A2"
widths = {"A": 12, "B": 12, "C": 10, "D": 18, "E": 18, "F": 34}
for col, width in widths.items():
    ws.column_dimensions[col].width = width
for row in range(2, ws.max_row + 1):
    ws.cell(row, 6).number_format = "0.0"
    for col in (2, 3, 4, 5):
        ws.cell(row, col).number_format = "#,##0"

chart = LineChart()
chart.title = "Six-term regulatory incidence per 100,000 words"
chart.y_axis.title = "Incidence per 100,000 words"
chart.x_axis.title = "Year"
chart.style = 13
chart.height = 10
chart.width = 22
data = Reference(ws, min_col=6, min_row=1, max_row=ws.max_row)
cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.legend = None
ws.add_chart(chart, "H2")

wb.save(output)
print(output)
