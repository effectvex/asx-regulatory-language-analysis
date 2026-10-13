"""Archive the histogram, thematic coding, and frame analysis before removal."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


def load_records(shard_directory: Path) -> list[dict]:
    records: list[dict] = []
    for path in sorted(shard_directory.glob("balanced-regulatory-shard-*.json")):
        records.extend(json.loads(path.read_text(encoding="utf-8-sig")))
    if len(records) != 1045:
        raise ValueError(f"Expected 1,045 records; found {len(records)}")
    return records


def load_metadata(path: Path) -> dict[tuple[str, int], dict]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return {
            (row["Ticker"], int(row["Year"])): row
            for row in csv.DictReader(handle)
        }


def style_sheet(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1769AA")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column[:100]]
        width = min(44, max(10, max(map(len, values), default=10) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def write_workbook(records: list[dict], metadata: dict, output: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    observations = workbook.create_sheet("Histogram Observations")
    observations.append([
        "Year", "Ticker", "Company", "Sector", "Subsector",
        "Extracted words", "Six-term occurrences", "Incidence per 100,000",
    ])
    rates_by_year: dict[int, list[tuple[str, str, float]]] = defaultdict(list)
    theme_by_year: dict[int, Counter] = defaultdict(Counter)
    frame_by_year: dict[int, Counter] = defaultdict(Counter)
    theme_total: Counter = Counter()
    frame_total: Counter = Counter()

    for record in sorted(records, key=lambda row: (int(row["Year"]), row["Ticker"])):
        year = int(record["Year"])
        key = (record["Ticker"], year)
        meta = metadata[key]
        occurrences = sum(int(value) for value in record["Terms"].values())
        words = int(record["Words"])
        incidence = 100000 * occurrences / words if words else 0.0
        observations.append([
            year, record["Ticker"], record["Company"], meta["Sector"],
            meta["Subsector"], words, occurrences, incidence,
        ])
        rates_by_year[year].append((record["Ticker"], record["Company"], incidence))
        theme_by_year[year].update({key: int(value) for key, value in record["Themes"].items()})
        frame_by_year[year].update({key: int(value) for key, value in record["Frames"].items()})
        theme_total.update({key: int(value) for key, value in record["Themes"].items()})
        frame_total.update({key: int(value) for key, value in record["Frames"].items()})
    style_sheet(observations)
    for cell in observations["H"][1:]:
        cell.number_format = "0.0"

    distribution = workbook.create_sheet("Yearly Distribution")
    distribution.append([
        "Year", "Companies", "Mean incidence", "Median incidence",
        "Standard deviation", "Minimum", "Maximum", "Three-sigma outliers",
    ])
    for year in sorted(rates_by_year):
        rows = rates_by_year[year]
        values = [row[2] for row in rows]
        mean = statistics.fmean(values)
        standard_deviation = statistics.pstdev(values)
        outliers = [
            f"{ticker} ({value:.1f})"
            for ticker, _, value in rows
            if standard_deviation and abs(value - mean) > 3 * standard_deviation
        ]
        distribution.append([
            year, len(values), mean, statistics.median(values), standard_deviation,
            min(values), max(values), "; ".join(outliers) or "None",
        ])
    style_sheet(distribution)
    for row in distribution.iter_rows(min_row=2, min_col=3, max_col=7):
        for cell in row:
            cell.number_format = "0.0"

    theme_summary = workbook.create_sheet("Theme Totals")
    theme_summary.append(["Theme", "Coded passages", "Share of theme assignments"])
    theme_denominator = sum(theme_total.values())
    for theme, count in theme_total.most_common():
        theme_summary.append([theme, count, count / theme_denominator if theme_denominator else 0])
    style_sheet(theme_summary)
    for cell in theme_summary["C"][1:]:
        cell.number_format = "0.0%"

    theme_year = workbook.create_sheet("Themes by Year")
    themes = list(records[0]["Themes"])
    theme_year.append(["Year", *themes])
    for year in sorted(theme_by_year):
        theme_year.append([year, *(theme_by_year[year][theme] for theme in themes)])
    style_sheet(theme_year)

    frame_summary = workbook.create_sheet("Frame Totals")
    frame_summary.append(["Frame", "Coded passages", "Share of coded passages"])
    frame_denominator = sum(frame_total.values())
    for frame, count in frame_total.most_common():
        frame_summary.append([frame, count, count / frame_denominator if frame_denominator else 0])
    style_sheet(frame_summary)
    for cell in frame_summary["C"][1:]:
        cell.number_format = "0.0%"

    frame_year = workbook.create_sheet("Frames by Year")
    frames = list(records[0]["Frames"])
    frame_year.append(["Year", *frames])
    for year in sorted(frame_by_year):
        frame_year.append([year, *(frame_by_year[year][frame] for frame in frames)])
    style_sheet(frame_year)

    notes = workbook.create_sheet("Methodology")
    notes.append(["Field", "Description"])
    notes.append(["Archive scope", "Histogram, thematic coding, and frame analysis removed from the main dashboard."])
    notes.append(["Histogram", "Company-year six-term incidence per 100,000 extracted words; yearly statistics use population standard deviation."])
    notes.append(["Themes", "Deterministic multi-label dictionary codes; one passage can receive multiple themes."])
    notes.append(["Frames", "Deterministic mutually exclusive primary frame assigned to each retained passage."])
    notes.append(["Validation", "Exploratory coding only; independent human validation has not been completed."])
    style_sheet(notes)

    theme_chart = BarChart()
    theme_chart.title = "Theme assignments"
    theme_chart.y_axis.title = "Coded passages"
    theme_chart.add_data(Reference(theme_summary, min_col=2, min_row=1, max_row=theme_summary.max_row), titles_from_data=True)
    theme_chart.set_categories(Reference(theme_summary, min_col=1, min_row=2, max_row=theme_summary.max_row))
    theme_chart.height = 8
    theme_chart.width = 18
    theme_summary.add_chart(theme_chart, "E2")

    frame_chart = BarChart()
    frame_chart.title = "Primary frame assignments"
    frame_chart.y_axis.title = "Coded passages"
    frame_chart.add_data(Reference(frame_summary, min_col=2, min_row=1, max_row=frame_summary.max_row), titles_from_data=True)
    frame_chart.set_categories(Reference(frame_summary, min_col=1, min_row=2, max_row=frame_summary.max_row))
    frame_chart.height = 7
    frame_chart.width = 16
    frame_summary.add_chart(frame_chart, "E2")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def write_html(source: Path, output: Path) -> None:
    html = source.read_text(encoding="utf-8-sig")
    head = re.search(r"(?s)<head>.*?</head>", html)
    sections = re.search(r"(?s)<h2>4\. Histogram.*?(?=<footer>)", html)
    script = re.search(r"(?s)<script>.*?</script>", html)
    if not (head and sections and script):
        raise ValueError("Could not locate the three removable sections in the source report")
    archived_head = re.sub(
        r"<title>.*?</title>",
        "<title>Archived histogram, thematic coding and frame analysis</title>",
        head.group(0),
    )
    archived = (
        "<!doctype html><html lang='en'>" + archived_head + "<body><main>"
        "<header><h1>Archived histogram, thematic coding and frame analysis</h1>"
        "<p class='dek'>Preserved from the balanced ASX200 regulatory-language report before these exploratory sections were removed from the main dashboard.</p>"
        "<div class='meta'><span class='pill'>95 companies</span><span class='pill'>1,045 reports</span><span class='pill'>2015-2025</span></div></header>"
        + sections.group(0)
        + "<footer>Archived analytical sections. Theme and frame coding is exploratory and has not received independent human validation.</footer>"
        "</main>" + script.group(0) + "</body></html>"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(archived, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-html", type=Path, required=True)
    parser.add_argument("--shards", type=Path, required=True)
    parser.add_argument("--metadata", type=Path, required=True)
    parser.add_argument("--output-directory", type=Path, required=True)
    args = parser.parse_args()

    records = load_records(args.shards)
    metadata = load_metadata(args.metadata)
    write_html(args.source_html, args.output_directory / "histogram-themes-frames-2015-2025.html")
    write_workbook(records, metadata, args.output_directory / "histogram-themes-frames-2015-2025.xlsx")
    print(f"written={args.output_directory.resolve()}")


if __name__ == "__main__":
    main()
